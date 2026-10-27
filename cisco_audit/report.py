"""Build one readable, sanitized audit sheet grouped by device IP."""
from collections import Counter
from datetime import datetime, timezone
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

from . import __version__
from .core import safe_text
from .checklist import evaluate


COLUMNS = (
    'IP / device', 'Section', 'Check / item', 'Sanitized result', 'Status',
    'Evidence ID', 'Read-only command', 'Scope / follow-up', 'Collected (UTC)',
    'Auditor decision', 'Owner', 'Reviewer comment', 'Review date',
)
WIDTHS = (23, 20, 35, 74, 18, 20, 28, 62, 24, 18, 18, 40, 16)
LAST_COLUMN = len(COLUMNS)
NAVY = '14283E'
NAVY_LIGHT = '203B56'
TEAL = '087E78'
PALE = 'EFF6F7'
WHITE = 'FFFFFF'
TEXT = '233447'
MUTED = '617386'


def _excel_value(value):
    """Escape formula prefixes and characters forbidden in Excel XML."""
    if not isinstance(value, str):
        return value
    return safe_text(re.sub(r'[\ufffe\uffff\ud800-\udfff]', '', value))


def _append(ws, values):
    ws.append([_excel_value(value) for value in values])
    return ws.max_row


def _meta(ws, row, left_label, left_value, right_label, right_value):
    ws.cell(row, 1, left_label)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.cell(row, 2, _excel_value(left_value))
    ws.cell(row, 5, right_label)
    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=LAST_COLUMN)
    ws.cell(row, 6, _excel_value(right_value))
    for column in (1, 5):
        cell = ws.cell(row, column)
        cell.font = Font(name='Aptos', bold=True, color=MUTED, size=10)
    for column in (2, 6):
        cell = ws.cell(row, column)
        cell.font = Font(name='Aptos', color=TEXT, size=10)
    ws.row_dimensions[row].height = 22


def _device_identity(result, evtext, platform):
    commands = ('show version', 'show inventory') if platform == 'cisco' else ('show system info',)
    identity = '\n'.join(evtext[c] for c in commands if c in evtext and not evtext[c].startswith('['))
    hostname = ''
    for line in identity.splitlines():
        if ' uptime is ' in line:
            hostname = line.split(' uptime is ', 1)[0].strip()
        elif line.lower().startswith('hostname:'):
            hostname = line.split(':', 1)[1].strip()
    return hostname, identity


def workbook(run):
    wb = Workbook()
    wb.properties.title = 'Network Access Review'
    wb.properties.subject = 'Read-only administrative access evidence'
    wb.properties.creator = 'Network Access Review'
    wb.properties.description = ('Sanitized, read-only device evidence. '
                                 'This workbook is not a compliance certification.')
    ws = wb.active
    ws.title = 'Audit Results'
    ws.sheet_properties.tabColor = TEAL
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'D12'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = '1:11'

    ws.merge_cells(f'A1:{get_column_letter(LAST_COLUMN)}1')
    ws['A1'] = 'NETWORK ACCESS REVIEW  /  AUDIT RESULTS'
    ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws['A1'].font = Font(name='Aptos Display', bold=True, size=18, color=WHITE)
    ws['A1'].alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[1].height = 43
    ws.merge_cells(f'A2:{get_column_letter(LAST_COLUMN)}2')
    ws['A2'] = 'One sheet · each device grouped by IP · sanitized, read-only collection'
    ws['A2'].font = Font(name='Aptos', size=10, color=TEAL, italic=True)
    ws['A2'].alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[2].height = 26

    platform = run.get('platform', 'cisco')
    results = run.get('results', [])
    _meta(ws, 4, 'Project', run.get('project_name', ''), 'Scan ID', run.get('id', ''))
    _meta(ws, 5, 'Platform', 'Cisco IOS / IOS XE' if platform == 'cisco' else 'Palo Alto PAN-OS', 'Started (UTC)', run.get('started', ''))
    _meta(ws, 6, 'Targets', len(results), 'Finished (UTC)', run.get('finished') or 'In progress')
    _meta(ws, 7, 'SSH profile', 'Legacy compatibility' if run.get('legacy') else 'Modern', 'Rerun of', run.get('rerun_of') or 'Original scan')
    counts = Counter(result.get('status', 'unknown') for result in results)
    _meta(ws, 8, 'Device status', ' · '.join(f'{status}: {count}' for status, count in sorted(counts.items())) or 'No devices', 'Evidence', 'Filtered output only')
    generated = datetime.now(timezone.utc).isoformat(timespec='seconds')
    _meta(ws, 10, 'Tool version', __version__, 'Report generated (UTC)', generated)
    ws.merge_cells(f'A9:{get_column_letter(LAST_COLUMN)}9')
    ws['A9'] = ('Observed means extracted, not compliant. Missing evidence does not prove a control is absent. '
                'Status legend: Observed = filtered field extracted; Attention = visible condition to review; '
                'Not Available / Not collected = not established from this snapshot; Manual verification = external records required. '
                'MFA enforcement requires independent provider-side verification.')
    ws['A9'].font = Font(name='Aptos', size=10, color=NAVY_LIGHT, italic=True)
    ws['A9'].fill = PatternFill('solid', fgColor=PALE)
    ws['A9'].alignment = Alignment(wrap_text=True, vertical='center', indent=1)
    ws.row_dimensions[9].height = 52

    for col, heading in enumerate(COLUMNS, 1):
        cell = ws.cell(11, col, heading)
        cell.font = Font(name='Aptos', bold=True, color=WHITE, size=10)
        cell.fill = PatternFill('solid', fgColor=NAVY_LIGHT)
        cell.alignment = Alignment(vertical='center', wrap_text=True, indent=1)
    ws.row_dimensions[11].height = 31
    for col, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    evidence_rows = {}
    evidence_references = []
    bottom_line = Border(bottom=Side(style='hair', color='D9E5E9'))
    for number, result in enumerate(results, 1):
        did = f'D{number:04d}'
        ip = str(result.get('ip', ''))
        port = result.get('port', 22)
        endpoint = ip if port == 22 else f'{ip}:{port}'
        device_platform = result.get('platform', platform)
        status = result.get('status', 'unknown')
        timestamp = result.get('time') or ''
        evidence = result.get('evidence', [])
        evmap = {ev['command']: f'{did}-E{index:02d}' for index, ev in enumerate(evidence, 1)}
        evtext = {ev['command']: ev['text'] for ev in evidence}
        hostname, identity = _device_identity(result, evtext, device_platform)

        band = ws.max_row + 2
        ws.merge_cells(start_row=band, start_column=1, end_row=band, end_column=LAST_COLUMN)
        title = f'{number:02d}  /  {endpoint}    ·    {hostname or "Hostname unavailable"}    ·    {device_platform.upper()}    ·    {status.upper()}'
        cell = ws.cell(band, 1, _excel_value(title))
        cell.fill = PatternFill('solid', fgColor=TEAL if status == 'completed' else NAVY_LIGHT)
        cell.font = Font(name='Aptos', color=WHITE, bold=True, size=12)
        cell.alignment = Alignment(vertical='center', indent=1)
        ws.row_dimensions[band].height = 32

        def add_row(section, check, value, state='', eid='', command='', note='', collected=timestamp):
            row = _append(ws, [endpoint, section, check, value, state, eid, command, note, collected, '', '', '', ''])
            ws.row_dimensions[row].height = 43 if section == 'Evidence' else 33
            for col in range(1, LAST_COLUMN + 1):
                c = ws.cell(row, col)
                c.font = Font(name='Aptos', size=10, color=TEXT, bold=(col == 2))
                c.alignment = Alignment(wrap_text=True, vertical='top', indent=1)
                c.border = bottom_line
                if row % 2 == 0:
                    c.fill = PatternFill('solid', fgColor='F7FAFB')
            if state in ('Observed', 'completed'):
                ws.cell(row, 5).fill = PatternFill('solid', fgColor='DDF1EA')
            elif state in ('Attention', 'partial', 'failed', 'host_key_unknown', 'interrupted'):
                ws.cell(row, 5).fill = PatternFill('solid', fgColor='FCE8D1')
            elif state in ('Not Available', 'Not collected', 'Manual verification', 'queued', 'connecting', 'collecting', 'cancelled'):
                ws.cell(row, 5).fill = PatternFill('solid', fgColor='E9EEF3')
            elif state == 'Needs review':
                ws.cell(row, 5).fill = PatternFill('solid', fgColor='E2E9FF')
            if eid and section != 'Evidence':
                evidence_references.append((ws.cell(row, 6), eid))
            return row

        completeness = 'Filtered output; not complete evidence' if status == 'completed' else 'Collection incomplete; inspect the issues below.'
        add_row('Device', 'Collection status', status, status, note=completeness)
        add_row('Device', 'Identity', identity or 'No identity data extracted.',
                'Observed' if identity else 'Not Available',
                eid=next((evmap[c] for c in (('show version', 'show inventory') if device_platform == 'cisco' else ('show system info',)) if c in evmap), ''),
                command='show version / show inventory' if device_platform == 'cisco' else 'show system info',
                note='Sanitized identity fields only; may be incomplete.', collected=timestamp if identity else '')

        for control in result.get('checklist') or evaluate(result, device_platform, did):
            linked = control.get('evidence_ids', [])
            remaining = ', '.join(linked[1:])
            note = control['follow_up']
            if remaining:
                note += ' Additional evidence: ' + remaining + '.'
            add_row('Checklist', control['id'] + ' · ' + control['area'],
                    control['observation'], control['status'],
                    linked[0] if linked else '', ', '.join(control.get('commands', [])),
                    note, timestamp if linked else '')

        account_source = 'show running-config' if device_platform == 'cisco' else 'show config running'
        facts_source = account_source
        accounts = result.get('accounts', [])
        facts = result.get('facts', [])
        fact_lines = [item.get('text', '') if isinstance(item, dict) else item for item in facts]
        account_lines = '\n'.join(
            f"{account.get('username', '')} — privilege/role: {account.get('privilege', 'unknown')}"
            for account in accounts
            if account.get('source', account_source) == account_source or account.get('role_source') == account_source
        )

        def add_check(check, source, value, limitation, attention=False, force_unavailable=False):
            available = bool(value and source in evmap and not evtext.get(source, '').startswith('[') and not value.startswith('['))
            state = 'Not Available' if force_unavailable else ('Attention' if attention else 'Observed') if available else 'Not Available'
            observed = value if available or force_unavailable else 'No data could be extracted from this command.'
            if len(observed) > 32000:
                limitation += ' Result truncated at the Excel cell limit; related evidence may also be truncated.'
            add_row('Check', check, observed, state,
                    evmap.get(source, '') if available or (force_unavailable and source in evmap) else '',
                    source, limitation,
                    timestamp if available or (force_unavailable and source in evmap) else '')

        if device_platform == 'cisco':
            add_check('Local accounts and privilege', account_source, account_lines,
                      'Extracted local accounts only; external AAA users and effective permissions are not established.')
            aaa = '\n'.join(line for line in fact_lines if line.lower().startswith('aaa '))
            add_check('AAA and method lists', facts_source, aaa,
                      'Visible configuration only; external server operation is not verified.')
            external = '\n'.join(line for line in fact_lines if re.match(r'(?i)^(?:aaa authentication|aaa authorization|aaa accounting|aaa group server|radius server|tacacs server|radius-server|tacacs-server)', line))
            mfa = ('External RADIUS/TACACS+ references observed:\n' + external + '\nProvider MFA enforcement was not verified.'
                   if external else 'No external RADIUS/TACACS+ reference was extracted; MFA enforcement cannot be determined.')
            add_check('MFA enforcement', facts_source, mfa,
                      'Device configuration does not prove a second-factor challenge or provider policy.', force_unavailable=True)
            vty = '\n'.join(line for line in fact_lines if line.lower().startswith('line vty '))
            attention = any(
                (' / transport input ' in line.lower()
                 and ('telnet' in line.lower() or line.lower().endswith(' all')))
                or line.lower().endswith(' / no login')
                for line in fact_lines)
            add_check('VTY access and transport', facts_source, vty,
                      'Filtered VTY fields; ACL definitions and all access paths may be incomplete.', attention=attention)
            ssh = evtext.get('show ip ssh', '')
            add_check('SSH status', 'show ip ssh', ssh,
                      'Other transport settings and access paths are not assessed.',
                      attention=ssh.lower().startswith('ssh disabled'))
        else:
            add_check('Administrators from configuration', account_source, account_lines,
                      'Visible configuration only; effective permissions and completeness are not established.')
            add_check('Administrator access configuration', facts_source, '\n'.join(fact_lines),
                      'Filtered management, authentication, lockout and profile settings; enforcement may differ.')
            auth = '\n'.join(line for line in fact_lines if line.lower().startswith(('admin ', 'authentication-profile ')))
            add_check('Administrator authentication and MFA', facts_source,
                      auth or 'No authentication-profile or per-admin authentication reference was extracted. MFA status is not established.',
                      'External RADIUS/SAML policy and actual MFA enforcement are not verified.', force_unavailable=True)
            listed = '\n'.join(account.get('username', '') for account in accounts if account.get('source') == 'show admins all')
            add_check('Administrators known to the device', 'show admins all', listed,
                      'Can include local, remote and Panorama admins; roles and effective access may be unknown.')

        for account in accounts:
            source = account.get('source', account_source)
            role_source = account.get('role_source', '' if source == 'show admins all' else source)
            note = 'Role source: ' + (role_source or 'not established') + '. Snapshot; external/effective access may differ.'
            add_row('Account', account.get('username', ''),
                    'Privilege / role: ' + str(account.get('privilege', 'unknown')),
                    'Observed', evmap.get(source, ''), source, note)
        for fact in facts:
            value = fact.get('text', '') if isinstance(fact, dict) else fact
            source = fact.get('source', facts_source) if isinstance(fact, dict) else facts_source
            add_row('Configuration', 'Access setting', value, 'Observed', evmap.get(source, ''), source,
                    'Filtered configuration field; actual enforcement is not proven.')
        for issue in result.get('issues', []):
            add_row('Issue', issue.get('command') or 'Collection', issue.get('message', ''),
                    'Attention', command=issue.get('command', ''), note='Data may be incomplete.', collected='')
        for index, ev in enumerate(evidence, 1):
            eid = f'{did}-E{index:02d}'
            text = ev.get('text', '')
            note = 'Sanitized excerpt; raw response is not exported.'
            if len(text) > 32000:
                note += ' Truncated at the Excel cell limit.'
            row = add_row('Evidence', ev.get('command', ''), text,
                          ev.get('status', 'filtered'), eid, ev.get('command', ''), note)
            evidence_rows[eid] = row
        if not evidence:
            add_row('Evidence', 'No filtered evidence available',
                    'Collection did not produce an exportable excerpt.', 'Not Available', collected='')

    for cell, eid in evidence_references:
        row = evidence_rows.get(eid)
        if row:
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'Audit Results'!A{row}", display=eid)
            cell.font = Font(name='Aptos', size=10, color=TEAL, underline='single')
    ws.print_area = f'A1:{get_column_letter(LAST_COLUMN)}{ws.max_row}'
    decision = DataValidation(type='list', formula1='"Pass,Fail,Needs review,Not available"', allow_blank=True)
    decision.errorTitle = 'Invalid decision'
    decision.error = 'Select Pass, Fail, Needs review, or Not available.'
    decision.promptTitle = 'Auditor decision'
    decision.prompt = 'Record the auditor decision for this row.'
    ws.add_data_validation(decision)
    decision.add(f'J12:J{max(ws.max_row, 12)}')
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
