import importlib.util
import sys
import types
import unittest
from io import BytesIO

HAS_DEPS = importlib.util.find_spec('openpyxl') is not None

# Report tests do not open SSH connections. If Paramiko is unavailable, a
# minimal import shim permits testing the actual XLSX generator independently;
# it does not test Paramiko, the SSH handshake, or real devices.
try:
    import paramiko  # noqa: F401
except ImportError:
    sys.modules['paramiko'] = types.ModuleType('paramiko')

@unittest.skipUnless(HAS_DEPS, 'openpyxl belum terpasang')
class WorkbookTests(unittest.TestCase):
    def load(self, platform, result):
        from openpyxl import load_workbook
        from cisco_audit.report import workbook
        run = {'id': 'test', 'platform': platform, 'started': '2026-01-01T00:00:00+00:00', 'results': [result]}
        return load_workbook(BytesIO(workbook(run)))

    def test_all_text_is_excel_safe_and_xml_compatible(self):
        from openpyxl import load_workbook
        from cisco_audit.report import workbook
        run = {'id': 'test', 'project_name': '=danger\ufffe', 'platform': 'cisco', 'started': '=formula', 'results': [{'ip': '=other', 'port': 22, 'status': 'failed', 'evidence': [], 'accounts': [], 'facts': [], 'issues': []}]}
        wb = load_workbook(BytesIO(workbook(run)))
        self.assertEqual(wb['Run Summary']['B2'].value, "'=danger")
        self.assertEqual(wb['Run Summary']['B4'].value, "'=formula")
        self.assertEqual(wb['Device Inventory']['C2'].value, "'=other")
        for ws in wb:
            self.assertFalse(any(c.data_type == 'f' for row in ws for c in row))

    def test_cisco_observations_are_evidence_linked_and_formula_safe(self):
        facts = ['line vty 0 4 / transport input telnet ssh', 'line vty 0 4 / ipv6 access-class MANAGEMENT in']
        r = {'ip': '192.0.2.1', 'port': 22, 'status': 'completed', 'evidence': [{'command': 'show running-config', 'text': '=malicious', 'status': 'filtered'}, {'command': 'show ip ssh', 'text': 'SSH Enabled - version 2.0', 'status': 'filtered'}], 'accounts': [{'username': 'admin', 'privilege': '15'}], 'facts': facts, 'issues': []}
        wb = self.load('cisco', r)
        self.assertEqual(wb.sheetnames, ['Run Summary', 'Device Inventory', 'Scan Results', 'Accounts', 'Access Configuration', 'Evidence', 'Collection Issues'])
        self.assertEqual(wb['Evidence']['E2'].value, "'=malicious")
        rows = {row[1]: row for row in list(wb['Scan Results'].values)[1:]}
        self.assertEqual(rows['Local accounts and privilege'][2], 'Observed')
        self.assertIn('admin', rows['Local accounts and privilege'][3])
        self.assertEqual(rows['Local accounts and privilege'][4], 'D0001-E01')
        self.assertEqual(rows['VTY access and transport'][2], 'Attention')
        self.assertIn('ipv6 access-class', rows['VTY access and transport'][3])
        self.assertEqual(rows['SSH status'][2], 'Observed')

    def test_large_evidence_is_explicitly_marked_truncated(self):
        text = 'x' * 33000
        result = {'ip': '192.0.2.7', 'port': 22, 'status': 'completed', 'evidence': [{'command': 'show ip ssh', 'text': text, 'status': 'filtered'}], 'accounts': [], 'facts': [], 'issues': []}
        wb = self.load('cisco', result)
        self.assertLessEqual(len(wb['Evidence']['E2'].value), 32000)
        self.assertIn('truncated', wb['Evidence']['F2'].value)
        row = next(row for row in list(wb['Scan Results'].values)[1:] if row[1] == 'SSH status')
        self.assertIn('truncated', row[7])
        self.assertIn('may also be truncated', row[7])

    def test_interrupted_scan_does_not_fabricate_evidence_timestamp(self):
        result = {'ip': '192.0.2.8', 'port': 22, 'status': 'interrupted',
                  'evidence': [{'command': 'show ip ssh', 'text': 'SSH Enabled - version 2.0', 'status': 'filtered'}],
                  'accounts': [], 'facts': [], 'issues': []}
        wb = self.load('cisco', result)
        self.assertIsNone(wb['Evidence']['D2'].value)
        self.assertIsNone(wb['Device Inventory']['G2'].value)
        row = next(row for row in list(wb['Scan Results'].values)[1:] if row[1] == 'SSH status')
        self.assertIsNone(row[6])
        self.assertEqual(row[2], 'Observed')

    def test_missing_evidence_is_never_observed_or_pass(self):
        wb = self.load('cisco', {'ip': '192.0.2.2', 'port': 22, 'status': 'failed', 'evidence': [], 'accounts': [], 'facts': [], 'issues': [{'command': '', 'message': 'Authentication failed'}]})
        rows = list(wb['Scan Results'].values)[1:]
        self.assertTrue(all(row[2] == 'Not Available' and row[4] is None for row in rows))
        self.assertEqual(wb['Collection Issues']['C2'].value, 'Authentication failed')

    def test_palo_sources_and_findings(self):
        r = {'ip': '192.0.2.3', 'port': 22, 'status': 'completed', 'platform': 'paloalto', 'evidence': [{'command': 'show system info', 'text': 'hostname: pa-test', 'status': 'filtered'}, {'command': 'show admins all', 'text': 'admin1', 'status': 'filtered'}, {'command': 'show config running', 'text': 'admin admin2 role superuser', 'status': 'filtered'}], 'accounts': [{'username':'admin1', 'privilege':'superuser', 'source':'show admins all', 'role_source':'show config running'}, {'username':'admin2', 'privilege':'superuser', 'source':'show config running'}], 'facts':['admin admin2 role superuser'], 'issues': []}
        wb = self.load('paloalto', r)
        self.assertEqual(wb['Device Inventory']['E2'].value, 'pa-test')
        rows = {row[1]: row for row in list(wb['Scan Results'].values)[1:]}
        self.assertEqual(rows['Administrators known to the device'][4], 'D0001-E02')
        self.assertIn('admin1', rows['Administrators known to the device'][3])
        self.assertEqual(rows['Administrators from configuration'][4], 'D0001-E03')
        self.assertIn('superuser', rows['Administrators from configuration'][3])
        self.assertEqual(wb['Accounts']['D2'].value, 'D0001-E02')
        self.assertEqual(wb['Accounts']['F2'].value, 'D0001-E03')
        self.assertNotIn('superuser', rows['Administrators known to the device'][3])
        self.assertIn('admin1', rows['Administrators from configuration'][3])

if __name__ == '__main__':
    unittest.main()
