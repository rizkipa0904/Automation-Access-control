"""End-to-end fixture matrix for all access-control checklist items.

These tests do not open network connections. They feed realistic, sanitized
device output through extraction, checklist evaluation, and workbook generation
to prove that every Cisco and PAN-OS control receives valid evidence/status and
that secrets never reach the report.
"""
import importlib.util
import sys
import textwrap
import types
import unittest
from io import BytesIO

try:
    import paramiko  # noqa: F401
except ImportError:
    sys.modules['paramiko'] = types.ModuleType('paramiko')

from cisco_audit.checklist import evaluate
from cisco_audit.core import (extract_config, extract_palo_admins, extract_palo_config,
                              safe_text, sanitize_show, validate_target)

HAS_REPORT = importlib.util.find_spec('openpyxl') is not None


CISCO_CONFIG = textwrap.dedent('''\
    service password-encryption
    no ip domain lookup
    clock timezone UTC 0 0
    aaa new-model
    aaa authentication login default group TACACS local
    aaa authentication enable default group TACACS enable
    aaa authorization console
    aaa authorization exec default group TACACS local
    aaa authorization commands 15 default group TACACS local
    aaa accounting exec default start-stop group TACACS
    aaa accounting commands 15 default start-stop group TACACS
    aaa group server tacacs+ TACACS
     server name TAC-1
    tacacs server TAC-1
     address ipv4 192.0.2.30
     key 7 TACACS_KEY_SECRET
    radius server RAD-1
     address ipv4 192.0.2.20 auth-port 1812 acct-port 1813
     key 7 RADIUS_KEY_SECRET
    radius-server host 192.0.2.21 auth-port 1812 acct-port 1813 key RADIUS_LEGACY_SECRET
    tacacs-server host 192.0.2.31 key TACACS_LEGACY_SECRET
    enable secret 9 ENABLE_SECRET_VALUE
    enable password 7 ENABLE_PASSWORD_VALUE
    username netops-ro privilege 5 secret 9 USER_SECRET_VALUE
    username auditor privilege 15 password 7 USER_PASSWORD_VALUE
    ip ssh version 2
    ip ssh time-out 60
    ip ssh authentication-retries 3
    ip ssh logging events
    ip ssh pubkey-chain
     username keyuser
      key-hash ssh-rsa SSH_KEY_HASH_SECRET
    line con 0
     login authentication CONSOLE
     exec-timeout 5 0
    line aux 0
     transport input none
    line vty 0 4
     login authentication default
     authorization exec default
     accounting commands 15 default
     transport input ssh
     access-class MGMT in
     ipv6 access-class MGMT6 in
     exec-timeout 10 0
     privilege level 15
    line vty 5 15
     login local
     transport input telnet ssh
    ip access-list extended MGMT
     permit tcp host 192.0.2.10 any eq 22
     deny ip any any log
    ip access-list standard MGMT
     permit 192.0.2.0 0.0.0.255
    ipv6 access-list MGMT6
     permit tcp host 2001:db8::10 any eq 22
    interface Vlan10
     description MANAGEMENT
     ip address 192.0.2.5 255.255.255.0
     ip vrf forwarding MGMT
    interface GigabitEthernet1/0/1
     no switchport
     ip address 192.0.2.6 255.255.255.0
    interface Loopback0
     ip address 192.0.2.7 255.255.255.255
    ip http server
    no ip http secure-server
    ip http authentication aaa
    ip http access-class MGMT
    snmp-server community SNMP_COMMUNITY_SECRET RO
    snmp-server user monitor SNMP_USER_SECRET
    logging trap warnings
    logging buffered 16384 warnings
    logging console warnings
    logging host 192.0.2.40 transport tcp port 514 vrf Mgmt
    logging host 192.0.2.41
    ntp authenticate
    ntp authentication-key 1 sha2 7 NTP_AUTH_SECRET
    ntp trusted-key 1
    ntp server vrf Mgmt 192.0.2.50 key 1 prefer
    ntp server 192.0.2.51
    login block-for 120 attempts 3 within 60
    login quiet-mode access-class MGMT
    login on-failure log
    service timestamps log datetime msec localtime
''')

CISCO_SECRETS = (
    'TACACS_KEY_SECRET', 'RADIUS_KEY_SECRET', 'RADIUS_LEGACY_SECRET',
    'TACACS_LEGACY_SECRET', 'ENABLE_SECRET_VALUE', 'ENABLE_PASSWORD_VALUE',
    'USER_SECRET_VALUE', 'USER_PASSWORD_VALUE', 'SSH_KEY_HASH_SECRET',
    'SNMP_COMMUNITY_SECRET', 'SNMP_USER_SECRET', 'NTP_AUTH_SECRET',
)


PALO_CONFIG = '''<config>
<mgt-config>
  <users>
    <entry name="fw-admin"><permissions><role-based><superuser>yes</superuser></role-based></permissions><authentication-profile>corp-radius</authentication-profile><phash>ADMIN_HASH_SECRET</phash></entry>
    <entry name="auditor"><permissions><role-based><custom><profile>readonly</profile></custom></role-based></permissions><authentication-profile>corp-saml</authentication-profile><password>ADMIN_PASSWORD_SECRET</password></entry>
  </users>
  <password-complexity><enabled>yes</enabled><minimum-length>12</minimum-length><minimum-uppercase-letters>1</minimum-uppercase-letters><minimum-lowercase-letters>1</minimum-lowercase-letters><minimum-numeric-letters>1</minimum-numeric-letters><minimum-special-characters>1</minimum-special-characters><password-history-count>5</password-history-count><expiration-period>90</expiration-period></password-complexity>
  <authentication-profile><entry name="corp-radius"><method><radius><server-profile>RADIUS_PROFILE_SECRET</server-profile></radius></method><factors><enable>yes</enable><factor><entry name="duo-mfa"/></factor></factors></entry></authentication-profile>
</mgt-config>
<shared>
  <authentication-profile><entry name="corp-saml"><method><saml><idp-profile>IDP_SECRET</idp-profile></saml></method></entry></authentication-profile>
  <admin-role><entry name="readonly"><role><device><cli>superreader</cli><xmlapi><enabled>yes</enabled></xmlapi><webui><monitor><logs>read-only</logs></monitor></webui></device></role><description>ROLE_DESC_SECRET</description></entry></admin-role>
</shared>
<devices><entry name="localhost.localdomain">
  <deviceconfig>
    <system>
      <permitted-ip><entry name="192.0.2.0/24"/><entry name="2001:db8::/32"/><entry name="BAD_SECRET_IP"/></permitted-ip>
      <service><disable-ssh>no</disable-ssh><disable-telnet>yes</disable-telnet><disable-http>yes</disable-http><disable-https>no</disable-https><disable-snmp>yes</disable-snmp></service>
      <timezone>UTC</timezone>
      <ntp-servers><primary-ntp-server><ntp-server-address>192.0.2.123</ntp-server-address><authentication-key>NTP_KEY_SECRET</authentication-key></primary-ntp-server></ntp-servers>
    </system>
    <setting><management><idle-timeout>30</idle-timeout><admin-lockout><failed-attempts>5</failed-attempts><lockout-time>15</lockout-time></admin-lockout></management></setting>
  </deviceconfig>
  <network>
    <profiles><interface-management-profile><entry name="MGMT"><ssh>yes</ssh><telnet>no</telnet><http>no</http><https>yes</https><snmp>no</snmp><permitted-ip><entry name="192.0.2.0/24"/></permitted-ip><password>PROFILE_PASSWORD_SECRET</password></entry></interface-management-profile></profiles>
    <interface><ethernet><entry name="ethernet1/1"><layer3><interface-management-profile>MGMT</interface-management-profile></layer3></entry></ethernet></interface>
  </network>
  <log-settings>
    <system><match-list><entry name="syslogs"><send-syslog><member>syslog-main</member></send-syslog></entry></match-list></system>
    <config><match-list><entry name="configs"><send-syslog><member>syslog-config</member></send-syslog></entry></match-list></config>
  </log-settings>
</entry></devices>
</config>'''

PALO_ADMINS = ('admin@fw> show admins all\n'
               'Admin   From        Client   Session-start      Idle-for\n'
               'fw-admin   103.0.0.5   CLI      06/06 15:06:09     00:00:00s\n'
               'auditor    -    API      -                  -\n'
               'panorama-admin   192.0.2.10   Web\n'
               '----   ----  ----   ----\n'
               'Total: 3\n'
               'admin@fw>')

PALO_SECRETS = (
    'ADMIN_HASH_SECRET', 'ADMIN_PASSWORD_SECRET', 'RADIUS_PROFILE_SECRET',
    'IDP_SECRET', 'ROLE_DESC_SECRET', 'BAD_SECRET_IP', 'NTP_KEY_SECRET',
    'PROFILE_PASSWORD_SECRET',
)


def cisco_result():
    facts, accounts = extract_config(CISCO_CONFIG)
    evidence = [
        {'command': 'show version', 'text': 'Cisco IOS version: 17.9.4', 'status': 'Filtered'},
        {'command': 'show running-config', 'text': 'filtered configuration evidence', 'status': 'Filtered'},
        {'command': 'show ip ssh', 'text': 'SSH Enabled - version 2.0', 'status': 'Filtered'},
        {'command': 'show privilege', 'text': 'Current privilege level is 15', 'status': 'Filtered'},
    ]
    return {'platform': 'cisco', 'ip': '192.0.2.1', 'port': 22, 'status': 'completed',
            'time': '2026-10-27T12:00:00+00:00', 'evidence': evidence,
            'accounts': accounts, 'facts': facts, 'issues': []}


def palo_result():
    facts, accounts = extract_palo_config(PALO_CONFIG)
    names = extract_palo_admins(PALO_ADMINS, 'admin@fw>')
    known = {account['username']: account for account in accounts}
    for name in names:
        known.setdefault(name, {'username': name, 'privilege': 'unknown', 'source': 'show admins all'})
    evidence = [
        {'command': 'show system info', 'text': 'hostname: fw-test\nmodel: PA-VM\nsw-version: 11.1.3', 'status': 'Filtered'},
        {'command': 'show admins all', 'text': '\n'.join(names), 'status': 'Filtered'},
        {'command': 'show config running', 'text': 'filtered configuration evidence', 'status': 'Filtered'},
    ]
    return {'platform': 'paloalto', 'ip': '192.0.2.2', 'port': 22, 'status': 'completed',
            'time': '2026-10-27T12:00:00+00:00', 'evidence': evidence,
            'accounts': list(known.values()), 'facts': facts, 'issues': []}


class AccessReviewMatrixTests(unittest.TestCase):
    def test_all_cisco_controls_receive_evidence_without_secret_leakage(self):
        result = cisco_result()
        rows = {row['id']: row for row in evaluate(result, 'cisco')}
        self.assertEqual(set(rows), {f'CS-{number:02d}' for number in range(1, 7)})
        self.assertEqual(rows['CS-06']['status'], 'Manual verification')
        for identifier in ('CS-01', 'CS-02', 'CS-03', 'CS-04', 'CS-05'):
            self.assertNotEqual(rows[identifier]['status'], 'Not collected', identifier)
            self.assertNotEqual(rows[identifier]['observation'], 'No matching device fields were extracted.', identifier)
        self.assertEqual(rows['CS-03']['status'], 'Attention')  # telnet on the second VTY range
        self.assertIn('line vty 0 4 / transport input ssh', rows['CS-02']['observation'])
        self.assertIn('ip access-list extended MGMT', rows['CS-03']['observation'])
        self.assertIn('ip http authentication aaa', rows['CS-04']['observation'])
        self.assertIn('logging host 192.0.2.40', rows['CS-05']['observation'])
        self.assertIn('Account: netops-ro', rows['CS-01']['observation'])
        self.assertIn('Account: auditor', rows['CS-06']['observation'])
        accounts = {account['username'] for account in result['accounts']}
        self.assertEqual(accounts, {'netops-ro', 'auditor'})  # pubkey-chain username is not an account
        rendered = repr(result)
        for secret in CISCO_SECRETS:
            self.assertNotIn(secret, rendered)

    def test_all_palo_controls_receive_evidence_without_secret_leakage(self):
        result = palo_result()
        rows = {row['id']: row for row in evaluate(result, 'paloalto')}
        self.assertEqual(set(rows), {f'PA-{number:02d}' for number in range(1, 7)})
        self.assertEqual(rows['PA-06']['status'], 'Manual verification')
        for identifier in ('PA-01', 'PA-02', 'PA-03', 'PA-04', 'PA-05'):
            self.assertNotEqual(rows[identifier]['status'], 'Not collected', identifier)
            self.assertNotEqual(rows[identifier]['observation'], 'No matching device fields were extracted.', identifier)
        self.assertIn('admin fw-admin role superuser', rows['PA-01']['observation'])
        self.assertIn('admin-role readonly cli superreader', rows['PA-01']['observation'])
        self.assertIn('password-complexity minimum-length 12', rows['PA-02']['observation'])
        self.assertIn('management permitted-ip 192.0.2.0/24', rows['PA-03']['observation'])
        self.assertIn('interface ethernet1/1 management-profile MGMT', rows['PA-03']['observation'])
        self.assertIn('management service disable-telnet yes', rows['PA-04']['observation'])
        self.assertIn('system log forwarding syslog-profile syslog-main', rows['PA-05']['observation'])
        self.assertIn('config log forwarding syslog-profile syslog-config', rows['PA-05']['observation'])
        self.assertIn('system ntp primary 192.0.2.123', rows['PA-05']['observation'])
        self.assertIn('Account: panorama-admin', rows['PA-06']['observation'])
        rendered = repr(result)
        for secret in PALO_SECRETS:
            self.assertNotIn(secret, rendered)

    def test_validate_target_accepts_the_supported_target_shapes(self):
        self.assertEqual(validate_target({'ip': '192.0.2.1', 'username': 'u', 'password': 'p'}),
                         {'ip': '192.0.2.1', 'port': 22, 'username': 'u', 'password': 'p'})
        self.assertEqual(validate_target({'ip': '2001:db8::1', 'port': '2222', 'username': 'u', 'password': 'p'})['port'], 2222)


class ParserRobustnessTests(unittest.TestCase):
    def test_extract_config_never_crashes_on_arbitrary_lines(self):
        samples = [
            '', '\x00\x01\x02', 'username', 'username  privilege', 'enable secret',
            'interface', 'line vty', 'ip access-list', 'banner motd ^C\nusername fake\n^C',
            'snmp-server community', 'logging host', 'ntp server', 'aaa',
            '😀\n\x1b[31mred\x1b[0m', 'a' * 50000, 'interface GigabitEthernet1/0/1\n ip address bad',
        ]
        for raw in samples:
            facts, accounts = extract_config(raw)
            self.assertIsInstance(facts, list)
            self.assertIsInstance(accounts, list)

    def test_banner_text_cannot_inject_accounts(self):
        raw = ('banner motd ^C\n'
               'username banner-fake privilege 15 secret 9 FAKE_SECRET\n'
               '^C\n'
               'username real-user privilege 7 secret 9 REAL_SECRET\n')
        facts, accounts = extract_config(raw)
        self.assertEqual(accounts, [{'username': 'real-user', 'privilege': '7'}])
        self.assertNotIn('banner-fake', repr((facts, accounts)))
        one_line, accounts = extract_config('banner login #username fake#\nusername real-user privilege 7\n')
        self.assertEqual(accounts, [{'username': 'real-user', 'privilege': '7'}])

    def test_palo_parsers_fail_safely_on_malformed_input(self):
        for raw in ('', 'not xml', '<config>', '<config><mgt-config>', '<config><mgt-config><users>'):
            with self.assertRaises(ValueError):
                extract_palo_config(raw)
        with self.assertRaises(ValueError):
            extract_palo_admins('unexpected output here', 'admin@fw>')
        self.assertEqual(extract_palo_admins('', 'admin@fw>'), [])

    def test_sanitizers_never_emit_arbitrary_secret_text(self):
        for command in ('show version', 'show inventory', 'show privilege', 'show ip ssh',
                        'show users', 'show ssh', 'unknown command'):
            text = sanitize_show(command, 'SECRET_MARKER \x00 \x1b[31m=cmd')
            self.assertIsInstance(text, str)
            self.assertNotIn('SECRET_MARKER', text)

    def test_safe_text_handles_control_and_bidi_characters(self):
        value = safe_text('\x00=1‮​')
        self.assertTrue(value.startswith("'"))
        self.assertNotIn('\x00', value)
        self.assertNotIn('‮', value)
        self.assertNotIn('​', value)


@unittest.skipUnless(HAS_REPORT, 'openpyxl is not installed')
class WorkbookMatrixTests(unittest.TestCase):
    def load(self, run):
        from openpyxl import load_workbook
        from cisco_audit.report import workbook
        book = load_workbook(BytesIO(workbook(run)))
        self.assertEqual(book.sheetnames, ['Audit Results'])
        return book.active

    def test_cisco_and_palo_workbooks_contain_all_controls_and_no_secrets(self):
        for platform, result, expected_ids, secrets in (
            ('cisco', cisco_result(), {f'CS-{number:02d}' for number in range(1, 7)}, CISCO_SECRETS),
            ('paloalto', palo_result(), {f'PA-{number:02d}' for number in range(1, 7)}, PALO_SECRETS),
        ):
            with self.subTest(platform=platform):
                run = {'id': 'matrix-scan', 'project_name': 'Matrix',
                       'platform': platform, 'started': '2026-10-27T12:00:00+00:00',
                       'finished': '2026-10-27T12:01:00+00:00', 'legacy': False,
                       'results': [result]}
                ws = self.load(run)
                checklist = [row for row in ws.values if row[1] == 'Checklist']
                self.assertEqual({row[2].split(' · ', 1)[0] for row in checklist}, expected_ids)
                self.assertEqual(len(checklist), 6)
                self.assertTrue(all(row[7] for row in checklist))
                account_rows = [row for row in ws.values if row[1] == 'Account']
                self.assertTrue(account_rows)
                evidence_rows = [row for row in ws.values if row[1] == 'Evidence']
                self.assertTrue(evidence_rows)
                self.assertEqual(ws.max_column, 13)
                self.assertEqual(ws['J11'].value, 'Auditor decision')
                self.assertFalse(any(cell.data_type == 'f' for row in ws for cell in row))
                rendered = '\n'.join(str(cell.value) for row in ws for cell in row if cell.value is not None)
                for secret in secrets:
                    self.assertNotIn(secret, rendered)
                links = [cell for row in ws for cell in row if cell.hyperlink]
                self.assertTrue(links)

    def test_workbook_handles_missing_optional_result_fields(self):
        ws = self.load({'id': 'minimal', 'results': [{'ip': '192.0.2.9', 'port': 22, 'status': 'completed'}]})
        self.assertEqual(len([row for row in ws.values if row[1] == 'Checklist']), 6)
        self.assertEqual(ws.max_column, 13)

    def test_workbook_handles_empty_run(self):
        ws = self.load({'id': 'empty', 'results': []})
        self.assertEqual(ws['A11'].value, 'IP / device')
        self.assertEqual(ws.max_column, 13)


if __name__ == '__main__':
    unittest.main()
